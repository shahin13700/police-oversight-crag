"""
tests/test_chunker.py
---------------------
Unit tests for src/ingestion/chunker.py.

We test:
1. chunk_document() produces chunks from a real or mock document
2. Each chunk has the required metadata fields populated
3. Section numbers are correctly extracted
4. Part names are correctly tracked across sections
5. Empty/structural paragraphs are excluded
6. The last chunk is not lost (boundary condition)

We use a synthetic Document built in memory so tests don't depend on the
real CSPA file being present (which is gitignored).

Branch: feature/ingestion-docx-chunker
Issue:  #3 — Section-aware legislative chunker
"""

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from src.ingestion.chunker import (
    chunk_document,
    extract_section_number,
    clean_part_name,
    LegislativeChunk,
)


# ---------------------------------------------------------------------------
# Helpers — build synthetic documents for testing
# ---------------------------------------------------------------------------

def add_paragraph_with_style(doc: Document, text: str, style_name: str) -> None:
    """
    Add a paragraph with a specific style to a Document.

    We add the style dynamically since our test document doesn't have
    the CSPA styles pre-loaded. This simulates the structure we found
    in the real CSPA .docx file.
    """
    para = doc.add_paragraph(text)
    # Try to apply the style — if it doesn't exist, create it
    try:
        para.style = doc.styles[style_name]
    except KeyError:
        # Style doesn't exist in this test doc — create a minimal one
        style = doc.styles.add_style(style_name, 1)  # 1 = paragraph style
        para.style = style


def build_sample_cspa_doc() -> Document:
    """
    Build a minimal synthetic document that mimics the CSPA .docx structure.

    This document has:
    - 1 part header
    - 2 sections, each with a headnote and body paragraphs
    - Some structural paragraphs that should be excluded
    """
    doc = Document()

    # Structural content — should be excluded from chunks
    add_paragraph_with_style(doc, "Community Safety and Policing Act, 2019", "shorttitle")
    add_paragraph_with_style(doc, "S.O. 2019, chapter 1", "chapter")
    add_paragraph_with_style(doc, "CONTENTS", "toc")

    # Part I header
    add_paragraph_with_style(doc, "PART I\nPrinciples and Interpretation", "partnum")

    # Section 1
    add_paragraph_with_style(doc, "Declaration of principles", "headnote")
    add_paragraph_with_style(
        doc,
        "1 Policing shall be provided throughout Ontario in accordance with the following principles.",
        "section"
    )
    add_paragraph_with_style(doc, "    1.    The need to ensure the safety of all persons.", "paragraph")
    add_paragraph_with_style(doc, "    2.    The importance of safeguarding fundamental rights.", "paragraph")

    # Section 2
    add_paragraph_with_style(doc, "Interpretation", "headnote")
    add_paragraph_with_style(doc, "2 (1)  In this Act,", "section")
    add_paragraph_with_style(doc, '"adequate and effective policing" means...', "definition")
    add_paragraph_with_style(doc, '"chief of police" means a chief of police...', "definition")
    add_paragraph_with_style(
        doc,
        "(2)  Words in the singular include the plural.",
        "subsection"
    )

    return doc


# ---------------------------------------------------------------------------
# Tests for extract_section_number()
# ---------------------------------------------------------------------------

class TestExtractSectionNumber:
    """Tests for the section number extraction helper."""

    def test_extracts_simple_section_number(self):
        """'1 Policing shall...' should return '1'"""
        result = extract_section_number("1 Policing shall be provided throughout Ontario.")
        assert result == "1"

    def test_extracts_section_with_subsection(self):
        """'2 (1)  In this Act,' should return '2(1)'"""
        result = extract_section_number("2 (1)  In this Act,")
        assert result == "2(1)"

    def test_extracts_double_digit_section(self):
        """'11 (1)  Adequate and effective policing...' should return '11(1)'"""
        result = extract_section_number("11 (1)  Adequate and effective policing means all of the following.")
        assert result == "11(1)"

    def test_returns_empty_for_subsection_only(self):
        """'(2)  For the purposes of...' has no leading section number."""
        result = extract_section_number("(2)  For the purposes of this section.")
        assert result == ""

    def test_returns_empty_for_non_section_text(self):
        """Definition text like '"adequate" means...' has no section number."""
        result = extract_section_number('"adequate and effective policing" means...')
        assert result == ""


# ---------------------------------------------------------------------------
# Tests for clean_part_name()
# ---------------------------------------------------------------------------

class TestCleanPartName:
    """Tests for the part name cleaning helper."""

    def test_cleans_multiline_part_name(self):
        """'PART III\\nRESPONSIBILITY...' should become 'PART III — RESPONSIBILITY...'"""
        result = clean_part_name("PART III\nRESPONSIBILITY FOR PROVIDING POLICING")
        assert result == "PART III — RESPONSIBILITY FOR PROVIDING POLICING"

    def test_cleans_simple_part_name(self):
        """Single-line part names should just be stripped."""
        result = clean_part_name("  PART I  ")
        assert result == "PART I"


# ---------------------------------------------------------------------------
# Tests for chunk_document()
# ---------------------------------------------------------------------------

class TestChunkDocument:
    """Tests for the main chunking function."""

    def test_produces_correct_number_of_chunks(self):
        """
        Our sample doc has 2 headnotes so should produce exactly 2 chunks.
        """
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="TEST_DOC")
        assert len(chunks) == 2

    def test_chunks_have_correct_type(self):
        """All returned items should be LegislativeChunk instances."""
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="TEST_DOC")
        for chunk in chunks:
            assert isinstance(chunk, LegislativeChunk)

    def test_section_titles_are_correct(self):
        """Chunk titles should match the headnote text."""
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="TEST_DOC")
        titles = [c.section_title for c in chunks]
        assert "Declaration of principles" in titles
        assert "Interpretation" in titles

    def test_section_numbers_are_extracted(self):
        """First chunk should have section number '1', second should have '2(1)'."""
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="TEST_DOC")

        # Find chunks by title to avoid ordering assumptions
        principles = next(c for c in chunks if c.section_title == "Declaration of principles")
        interpretation = next(c for c in chunks if c.section_title == "Interpretation")

        assert principles.section_number == "1"
        assert interpretation.section_number == "2(1)"

    def test_part_name_is_tracked(self):
        """Both chunks should have the correct part name."""
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="TEST_DOC")
        for chunk in chunks:
            assert "PART I" in chunk.part_name

    def test_source_doc_is_set(self):
        """source_doc should be passed through to all chunks."""
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="CSPA_2019")
        for chunk in chunks:
            assert chunk.source_doc == "CSPA_2019"

    def test_chunk_text_contains_section_content(self):
        """Chunk text should include both the title and body paragraphs."""
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="TEST_DOC")
        principles = next(c for c in chunks if c.section_title == "Declaration of principles")

        # Should contain the headnote
        assert "Declaration of principles" in principles.text
        # Should contain the section body
        assert "Policing shall be provided" in principles.text

    def test_structural_paragraphs_excluded(self):
        """shorttitle, chapter, toc styles should not appear in any chunk."""
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="TEST_DOC")
        all_text = " ".join(c.text for c in chunks)

        assert "S.O. 2019, chapter 1" not in all_text
        assert "CONTENTS" not in all_text

    def test_citation_format(self):
        """citation() should return properly formatted CSPA citation."""
        doc = build_sample_cspa_doc()
        chunks = chunk_document(doc, source_doc="CSPA_2019")
        principles = next(c for c in chunks if c.section_title == "Declaration of principles")
        assert principles.citation() == "CSPA s.1"

    def test_empty_document_returns_no_chunks(self):
        """An empty document should return an empty list, not crash."""
        doc = Document()
        chunks = chunk_document(doc, source_doc="EMPTY")
        assert chunks == []
